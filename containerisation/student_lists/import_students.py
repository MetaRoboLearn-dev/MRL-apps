#!/usr/bin/env python3
"""
Import students from an .xlsx file into the MRL database.

Usage:
    DATABASE_URL=... python import_students.py <students.xlsx>

    Or place a .env file in the containerisation/ folder (one level up).

Output:
    <input_basename>_with_passwords.xlsx  - written next to the input file
    Console summary with flagged duplicates (ime+prezime already in DB)

Password format : <1st letter of ime><prezime><5 random digits><random special char>
                  e.g.  "ITerzic48321?"
Username format : ime.prezime  (lowercase, Croatian diacritics normalised to ASCII)
                  e.g.  "ivan.terzic"
Actor (created_by): user with username "IvanTerzic1"
"""

import os
import sys
import random
import string
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse, urlunparse

import bcrypt
import openpyxl
import psycopg2
import psycopg2.errors

try:
    from dotenv import load_dotenv
    # Try .env one directory up (containerisation/.env)
    load_dotenv(Path(__file__).resolve().parent.parent / ".env")
except ImportError:
    pass  # dotenv optional; set DATABASE_URL directly in the environment


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

ACTOR_USERNAME = "IvanTerzic1"
STUDENT_ROLE   = "student"
SPECIAL_CHARS  = list("+-!?#$%")

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

_CROATIAN_MAP = str.maketrans({
    "\u010d": "c", "\u0107": "c", "\u0161": "s", "\u017e": "z", "\u0111": "d",
    "\u010c": "C", "\u0106": "C", "\u0160": "S", "\u017d": "Z", "\u0110": "D",
})


def normalize(text: str) -> str:
    """Replace Croatian diacritics with ASCII equivalents."""
    return text.translate(_CROATIAN_MAP)


def make_username(ime: str, prezime: str) -> str:
    return f"{normalize(ime.strip()).lower()}.{normalize(prezime.strip()).lower()}"


def make_password(ime: str, prezime: str) -> str:
    """<1st letter of ime><Prezime><5 random digits><random special char>"""
    first   = ime.strip()[0].upper()
    last    = normalize(prezime.strip())
    last    = last[0].upper() + last[1:].lower() if last else last
    digits  = "".join(random.choices(string.digits, k=5))
    special = random.choice(SPECIAL_CHARS)
    return f"{first}{last}{digits}{special}"


def hash_password(plain: str) -> str:
    return bcrypt.hashpw(plain.encode(), bcrypt.gensalt()).decode()


def db_connect(url: str):
    """Strip SQLAlchemy driver prefix, force host to localhost, and connect."""
    url = url.replace("postgresql+psycopg2://", "postgresql://")
    url = url.replace("postgresql+psycopg://",  "postgresql://")
    parsed = urlparse(url)
    userinfo = f"{parsed.username}:{parsed.password}@" if parsed.username else ""
    port     = f":{parsed.port}" if parsed.port else ""
    new_netloc = f"{userinfo}localhost{port}"
    url = urlunparse(parsed._replace(netloc=new_netloc))
    return psycopg2.connect(url)


def find_col(headers_lower: list, name: str) -> int:
    try:
        return headers_lower.index(name.lower())
    except ValueError:
        return -1


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    if len(sys.argv) < 2:
        sys.exit("Usage: python import_students.py <students.xlsx>")

    xlsx_path = Path(sys.argv[1])
    if not xlsx_path.exists():
        sys.exit(f"File not found: {xlsx_path}")

    db_url = os.environ.get("DATABASE_URL")
    if not db_url:
        sys.exit("DATABASE_URL environment variable is not set.")

    # -- Connect -------------------------------------------------------------
    try:
        conn = db_connect(db_url)
    except Exception as exc:
        sys.exit(f"Cannot connect to database: {exc}")

    cur = conn.cursor()

    cur.execute("SELECT id FROM users WHERE username = %s", (ACTOR_USERNAME,))
    row = cur.fetchone()
    if not row:
        cur.close(); conn.close()
        sys.exit(f"Actor user '{ACTOR_USERNAME}' not found in the database.")
    actor_id = row[0]

    cur.execute("SELECT id FROM roles WHERE name = %s", (STUDENT_ROLE,))
    row = cur.fetchone()
    if not row:
        cur.close(); conn.close()
        sys.exit(f"Role '{STUDENT_ROLE}' not found in the database.")
    student_role_id = row[0]

    # -- Read input xlsx -----------------------------------------------------
    wb_in = openpyxl.load_workbook(xlsx_path)
    ws_in = wb_in.active

    raw_headers   = [c.value for c in next(ws_in.iter_rows(min_row=1, max_row=1))]
    headers_lower = [str(h).strip().lower() if h is not None else "" for h in raw_headers]

    pii_col   = find_col(headers_lower, "prezime i ime")
    kod_col   = find_col(headers_lower, "kod")
    skola_col = find_col(headers_lower, "\u0161kola")  # škola
    if skola_col < 0:
        skola_col = find_col(headers_lower, "skola")   # fallback without diacritic

    if pii_col < 0:
        cur.close(); conn.close()
        sys.exit(
            "Could not find 'Prezime i ime' column.\n"
            f"Detected headers: {raw_headers}"
        )

    now = datetime.now(timezone.utc)

    imported   = []   # list of dicts
    duplicates = []   # (ime, prezime, existing_username)
    errors     = []   # (ime, prezime, reason)

    # -- Process rows --------------------------------------------------------
    for row in ws_in.iter_rows(min_row=2, values_only=True):
        if row[pii_col] is None:
            continue

        parts   = str(row[pii_col]).strip().split(" ", 1)
        if len(parts) < 2:
            continue
        prezime, ime = parts[0].strip(), parts[1].strip()
        kod   = str(row[kod_col]).strip()   if kod_col   >= 0 and row[kod_col]   is not None else ""
        skola = str(row[skola_col]).strip() if skola_col >= 0 and row[skola_col] is not None else ""

        if not ime or not prezime:
            continue

        # Duplicate check: same first+last name already in DB
        cur.execute(
            "SELECT username FROM users "
            "WHERE LOWER(first_name) = LOWER(%s) AND LOWER(last_name) = LOWER(%s)",
            (ime, prezime),
        )
        existing = cur.fetchone()
        if existing:
            duplicates.append((ime, prezime, existing[0]))
            continue

        base_username = make_username(ime, prezime)
        username      = base_username
        password      = make_password(ime, prezime)
        password_hash = hash_password(password)

        # Resolve username collisions with a numeric suffix
        suffix = 1
        while True:
            cur.execute("SELECT 1 FROM users WHERE username = %s", (username,))
            if not cur.fetchone():
                break
            username = f"{base_username}{suffix}"
            suffix  += 1

        try:
            cur.execute(
                """
                INSERT INTO users
                    (username, password_hash, first_name, last_name,
                     role_id, created_at, updated_at, created_by, updated_by, active)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, TRUE)
                """,
                (username, password_hash, ime, prezime,
                 student_role_id, now, now, actor_id, actor_id),
            )
            conn.commit()
            imported.append({
                "ime": ime, "prezime": prezime,
                "kod": kod, "skola": skola,
                "username": username, "password": password,
            })
        except psycopg2.errors.UniqueViolation:
            conn.rollback()
            errors.append((ime, prezime, "Username conflict after retry"))
        except Exception as exc:
            conn.rollback()
            errors.append((ime, prezime, str(exc)))

    cur.close()
    conn.close()

    # -- Write output xlsx ---------------------------------------------------
    out_path = xlsx_path.parent / (xlsx_path.stem + "_with_passwords.xlsx")
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Studenti"

    ws_out.append(["Ime", "Prezime", "\u0160kola", "KOD", "Username", "Lozinka"])
    for s in imported:
        ws_out.append([s["ime"], s["prezime"], s["skola"], s["kod"], s["username"], s["password"]])

    for col in ws_out.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=0)
        ws_out.column_dimensions[col[0].column_letter].width = max_len + 4

    wb_out.save(out_path)

    # -- Console summary -----------------------------------------------------
    sep = "=" * 50
    print(f"\n{sep}")
    print("  Import Complete")
    print(sep)
    print(f"  Imported  : {len(imported)}")
    print(f"  Duplicates: {len(duplicates)}")
    print(f"  Errors    : {len(errors)}")

    if duplicates:
        print(f"\n{'─' * 50}")
        print("  DUPLICATES  (ime+prezime already exists in DB)")
        print(f"{'─' * 50}")
        for ime, prezime, uname in duplicates:
            print(f"  !! {ime} {prezime}  ->  existing username: {uname}")

    if errors:
        print(f"\n{'─' * 50}")
        print("  ERRORS")
        print(f"{'─' * 50}")
        for ime, prezime, reason in errors:
            print(f"  XX {ime} {prezime}  ->  {reason}")

    print(f"\n  Output: {out_path}\n")


if __name__ == "__main__":
    main()
